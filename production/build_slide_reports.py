#!/usr/bin/env python3
"""Documents de livraison dérivés de la même source que le PPTX."""
from pathlib import Path
import sys
from datetime import date
ROOT=Path(__file__).resolve().parent.parent
sys.path.insert(0,str(ROOT))
from slides.content import SL,CH,PARTS,RANGES
from pptx import Presentation

p=Presentation(ROOT/'BELIVE_MONEY_formation_belive_money.pptx')
assert len(p.slides._sldIdLst)==len(SL)
for slide,s in zip(p.slides,SL):
 assert s.id in slide.notes_slide.notes_text_frame.text

head='''# BELIVE MONEY — storyboard de présentation

**Source éditoriale réellement disponible :** `BELIVE_MONEY_manuscrit_illustre.docx`, édition illustrée **adaptée et condensée** ; son PDF de lecture distinct comporte 69 pages. Le manuscrit intégral annoncé n'est pas présent. Le DOCX prévaut pour le contenu. Ce storyboard n'est pas une transcription du manuscrit intégral.

**Format :** 16:9, PowerPoint éditable ; **84 diapositives**, **179 minutes estimées** (2 h 59, pauses non incluses). Les notes intégrées dans le PPTX et le fichier de notes associé portent les explications et prompts complets. Une idée principale par slide ; les slides de chapitre montrent une méthode, les ateliers une action, les synthèses un rappel. Durées indicatives, ni seuil ni promesse. Les faits évolutifs sont séparés en bloc de vérification.

**Plan :** cadrage S001–S003 ; parties 1 à 8 S004–S066 ; prompts vérifiés du livre S067–S073 ; bloc de vérification S074–S079 ; mise en œuvre S080–S084. Les huit parties du livre gardent leurs chapitres dans leur ordre ; les trois blocs transversaux sont explicitement séparés et renvoient aux chapitres.

**Légende :** « texte visible » donne le contenu à l'écran et non les pieds de page/repères de partie ; « notes » renvoie aux notes complètes également reproduites dans `BELIVE_MONEY_notes_presentateur.md`. « Source/Crédits » indique les textes documentés et les schémas originaux ; les URL officielles sont dans le registre de sources. Les titres ne sont pas des intitulés de plateformes officielles.

'''
# Correct sequence: sections 67, prompts 68-73, section 74, factual 75-77, decision 78, section 79, roadmaps 80-81 checklist 82 etc.
head=head.replace('prompts vérifiés du livre S067–S073 ; bloc de vérification S074–S079 ; mise en œuvre S080–S084', 'prompts vérifiés du livre S067–S073 ; bloc de vérification S074–S078 ; mise en œuvre S079–S084')

def visible(s):
 text=[]
 if s.subtitle:text.append(s.subtitle)
 if s.nodes:text.extend(s.nodes)
 if s.exercise:text.append('Consigne : '+s.exercise)
 if s.kind=='cover':text.append('Formation bâtie sur le manuscrit illustré adapté (DOCX) disponible dans le dépôt. Repères évolutifs revérifiés sur les pages officielles le 28 septembre 2026.')
 if s.kind=='closing':text.append('Une compétence · Un public · Un contenu pilote · Une mesure · Une décision')
 return text

visuals={'cover':'Typographie et repère vertical, sans photographie ni logo.',
 'section':'Intertitre typographique, bandeau de couleur distinct par partie.',
 'manifesto':'Trois cartes éditables : engagements, pas promesses.',
 'path':'Huit étapes en deux rangées de quatre cartes ; flèches de gauche à droite, numéros 01–08.',
 'chapter':'Étapes numérotées verticales en formes natives, couleur de partie + numéros ; sous-titre = objectif.',
 'exercise':'Carte de consigne unique ; bandeau rappelant le statut simulé des exemples.',
 'recap':'Étapes numérotées verticales reprenant les objectifs du groupe de chapitres.',
 'prompt':'Quatre cartes de variables ; prompt intégral en notes et non en taille minuscule.',
 'facts':'Quatre cartes verticales de conditions à recontrôler ; source officielle en pied.',
 'decision':'Matrice 2 × 2 de décisions locales ; lecture ligne par ligne.',
 'roadmap':'Quatre cartes chronologiques avec flèches ; calendrier indicatif.',
 'checklist':'Cinq lignes avec cases à cocher vides.',
 'closing':'Cinq cartes d’action et une phrase à remplir.',
 'resources':'Quatre entrées sources ; renvoi au registre URL.'}
story=[head]
notes=['# BELIVE MONEY — notes du présentateur\n\nReprise intégrale des notes présentes dans le PPTX. Durée totale indicative : **179 minutes** hors pauses. Les exemples cités par le livre sont des simulations pédagogiques, non des témoignages. Les prompts sont ceux des rubriques « PROMPT COPIABLE » du DOCX ; ne pas utiliser les variables sans contrôle.\n\n']
plan=['# BELIVE MONEY — plan des illustrations des diapositives\n\nLe diaporama ne contient **aucune photographie, image IA, logo ou capture** ; il utilise 100 % de formes, flèches, numéros, cartes et texte éditables PowerPoint, sans fichier image tiers et sans mention d’une charte officielle. Une image décorative n’apporterait rien aux 39 chapitres : chaque chapitre a au contraire son schéma pédagogique natif (étapes de la méthode) au format 16:9. Chaque forme porte une description OOXML (`descr`) ; le sens est également transmis par textes/numéros et ordre. Les seules flèches indiquent la lecture explicitement nommée, jamais une relation causale chiffrée. Les détails restent dans les notes.\n\n**Palette proposée, non officielle :** bleu nuit #081022, panneaux #121E38, texte #F2F6FC, accent cyan/bleu/violet/vert selon partie. Les plateformes ne sont pas représentées par leurs logos.\n\n']
for s in SL:
 part=(f'Partie {s.part:02d}' if 1<=s.part<=8 else f'Bloc transversal {s.part-8}' if s.part>=9 else 'Ouverture / clôture')
 node_text=' ; '.join(visible(s)) or '(titre et repère de partie uniquement)'
 n=s.notes.strip()
 story.append(f'''## {s.id} — {s.title}

- **Partie / emplacement :** {part} ; slide {int(s.id[1:]):02d}/84.
- **Titre-message :** {s.title}.
- **Objectif pédagogique :** {s.subtitle or (CH[s.chapter]['goal'] if s.chapter else 'Faire progresser le parcours de formation.')}
- **Idée unique :** {'Décider et réaliser la consigne de l’atelier.' if s.kind=='exercise' else 'Passer à une nouvelle partie.' if s.kind=='section' else 'Conserver un seul repère de méthode vérifiable.'}
- **Texte visible :** {node_text}
- **Exemple / exercice :** {s.exercise or (CH[s.chapter]['example'] if s.chapter and s.kind in ('chapter','recap') else 'Question orale : quelle décision pouvez-vous documenter maintenant ?' if s.kind=='recap' else 'Pas d’exercice supplémentaire.')}
- **Notes :** {n}
- **Visuel et emplacement :** {visuals[s.kind]} Positions : titre en haut, corps central, référence et minutage en pied ; emplacement précis détaillé dans le plan visuel.
- **Chapitre / partie source :** {s.source}.
- **Sources / crédits :** {s.slidesource} ; schéma éditable original BELIVE MONEY, construit pour cette formation. Voir registre `BELIVE_MONEY_credits_sources_images.md` pour [T1]–[G1] et réserves.
- **Durée :** {s.duration} min.
''')
 notes.append(f'## {s.id} · {s.title} ({s.duration} min)\n\n{n}\n\n**Source :** {s.source}.\n\n')
 plan.append(f'- **{s.id} — {s.title} :** {visuals[s.kind]} **Placement :** bandeau de partie et titre dans le tiers haut ; contenu au centre ; source et minuterie en bas. **Texte alternatif :** chacun des textes et formes logiques possède une description explicite dans le PPTX ; source sonore : aucune. **Crédit :** composition originale éditable, dérivée de {s.source}.\n')
(ROOT/'BELIVE_MONEY_storyboard_presentation.md').write_text(''.join(story),encoding='utf-8')
(ROOT/'BELIVE_MONEY_notes_presentateur.md').write_text(''.join(notes),encoding='utf-8')
(ROOT/'BELIVE_MONEY_plan_illustrations_slides.md').write_text(''.join(plan),encoding='utf-8')
print('rapports :',len(SL),'slides ;',sum(s.duration for s in SL),'min')
