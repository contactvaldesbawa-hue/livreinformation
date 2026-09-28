#!/usr/bin/env python3
"""Aperçu interne de contrôle du PPTX livré.

Lit le fichier PowerPoint réel (positions, couleurs, tailles de police,
textes) et produit une image par diapositive. Le rendu est approximatif ;
il ne remplace pas une ouverture PowerPoint. Aucun PDF n'est exporté.
"""
from pathlib import Path
import sys

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from slides.metrics import wrap_lines  # noqa: E402

PPTX = ROOT / "BELIVE_MONEY_formation_belive_money.pptx"
OUTDIR = ROOT / "BELIVE_MONEY_apercus_slides"
OUTDIR.mkdir(exist_ok=True)

SCALE = 2.0  # pixels par point
EMU_PT = 12700.0
prs = Presentation(PPTX)
SW = prs.slide_width / EMU_PT
SH = prs.slide_height / EMU_PT
PXW, PXH = int(SW * SCALE), int(SH * SCALE)

FB = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
FR = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
_fc = {}


def font(bold, pt):
    key = (bold, round(pt, 1))
    if key not in _fc:
        _fc[key] = ImageFont.truetype(FB if bold else FR, max(6, int(round(pt * SCALE))))
    return _fc[key]


def color_of(c):
    return (c[0], c[1], c[2])


def prst(shape):
    try:
        g = shape._element.spPr.find(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom")
        return g.get("prst") if g is not None else None
    except Exception:
        return None


def draw_shape(d, sh):
    l = sh.left / EMU_PT * SCALE
    t = sh.top / EMU_PT * SCALE
    w = max(1, sh.width / EMU_PT * SCALE)
    h = max(1, sh.height / EMU_PT * SCALE)
    geom = prst(sh)
    fill = None
    try:
        if sh.fill.type == MSO_FILL.SOLID:
            fill = color_of(sh.fill.fore_color.rgb)
    except Exception:
        fill = None
    line = None
    try:
        if sh.line.fill.type == MSO_FILL.SOLID:
            line = color_of(sh.line.color.rgb)
    except Exception:
        line = None
    if fill is None and line is None:
        return l, t, w, h
    if geom == "ellipse":
        d.ellipse([l, t, l + w, t + h], fill=fill, outline=line)
    elif geom == "rightArrow":
        body = h * 0.5
        d.polygon([(l, t + h / 2 - body / 2), (l + w * 0.55, t + h / 2 - body / 2),
                   (l + w * 0.55, t), (l + w, t + h / 2),
                   (l + w * 0.55, t + h), (l + w * 0.55, t + h / 2 + body / 2),
                   (l, t + h / 2 + body / 2)], fill=fill, outline=fill)
    elif geom == "roundRect":
        try:
            adj = sh.adjustments[0]
        except Exception:
            adj = 0.16667
        r = adj * min(w, h)
        d.rounded_rectangle([l, t, l + w, t + h], radius=r, fill=fill, outline=line,
                            width=int(1.2 * SCALE))
    else:
        d.rectangle([l, t, l + w, t + h], fill=fill, outline=line, width=int(1.2 * SCALE))
    return l, t, w, h


def draw_text(d, sh, l, t, w, h):
    tf = sh.text_frame
    paras = []
    for p in tf.paragraphs:
        txt = "".join(r.text for r in p.runs)
        if not p.runs:
            paras.append(("", 18, False, (255, 255, 255), p.alignment))
            continue
        size = max((r.font.size.pt for r in p.runs if r.font.size), default=18)
        bold = any(r.font.bold for r in p.runs)
        try:
            col = color_of(p.runs[0].font.color.rgb)
        except Exception:
            col = (255, 255, 255)
        paras.append((txt, size, bold, col, p.alignment))
    # hauteur totale
    blocks = []
    for txt, size, bold, col, align in paras:
        lines = wrap_lines(txt, size, w / SCALE, bold) if txt else [""]
        blocks.append((lines, size, bold, col, align))
    total = sum(len(b[0]) * b[1] * 1.14 * SCALE for b in blocks)
    anchor = tf.vertical_anchor
    y = t
    if anchor == MSO_ANCHOR.MIDDLE:
        y = t + (h - total) / 2
    for lines, size, bold, col, align in blocks:
        for ln in lines:
            f = font(bold, size)
            tw = d.textlength(ln, font=f)
            x = l
            if align == PP_ALIGN.CENTER:
                x = l + (w - tw) / 2
            elif align == PP_ALIGN.RIGHT:
                x = l + w - tw
            d.text((x, y), ln, font=f, fill=col)
            y += size * 1.14 * SCALE


frames = []
for i, slide in enumerate(prs.slides, 1):
    img = Image.new("RGB", (PXW, PXH), (8, 16, 34))
    d = ImageDraw.Draw(img)
    for sh in slide.shapes:
        l, t, w, h = draw_shape(d, sh)
        if sh.has_text_frame and sh.text_frame.text.strip():
            draw_text(d, sh, l, t, w, h)
    d.rectangle([0, 0, PXW, PXH], outline=(120, 140, 180), width=1)
    img = img.resize((PXW // 2, PXH // 2), Image.LANCZOS)
    p = OUTDIR / f"slide_{i:03d}.png"
    img.save(p)
    frames.append(p)

print(f"{len(frames)} aperçus PNG dans {OUTDIR.name}/ (contrôle interne seulement ; aucun PDF exporté)")
