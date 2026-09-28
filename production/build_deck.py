#!/usr/bin/env python3
"""Construit le diaporama PPTX natif BELIVE MONEY à partir du storyboard.

Aucune image externe, aucun logo, aucune capture : toutes les compositions
sont des formes et du texte natifs (donc éditables dans PowerPoint).
Chaque diapositive reçoit : titre, corps, texte alternatif, source et notes.
"""
from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from slides import content as C  # noqa: E402
from slides.metrics import fit_size  # noqa: E402

OUT = ROOT / "BELIVE_MONEY_formation_belive_money.pptx"

# ---------------------------------------------------------------- direction
BG = RGBColor(0x08, 0x10, 0x22)
PANEL = RGBColor(0x12, 0x1E, 0x38)
PANEL2 = RGBColor(0x1B, 0x2A, 0x4A)
INK = RGBColor(0xF2, 0xF6, 0xFC)
MUTED = RGBColor(0x9F, 0xB0, 0xCC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

ACCENTS = {
    1: RGBColor(0x22, 0xD3, 0xEE), 2: RGBColor(0x38, 0xBD, 0xF8),
    3: RGBColor(0x81, 0x8C, 0xF8), 4: RGBColor(0xA7, 0x8B, 0xFA),
    5: RGBColor(0xC0, 0x84, 0xFC), 6: RGBColor(0x2D, 0xD4, 0xBF),
    7: RGBColor(0x34, 0xD3, 0x99), 8: RGBColor(0x60, 0xA5, 0xFA),
    9: RGBColor(0x94, 0xA3, 0xB8), 10: RGBColor(0x94, 0xA3, 0xB8),
    11: RGBColor(0x94, 0xA3, 0xB8), 0: RGBColor(0x22, 0xD3, 0xEE),
}
FONT = "Arial"
SW, SH = 13.333, 7.5
W_PT = SW * 72.0


def accent(part: int) -> RGBColor:
    return ACCENTS.get(part, ACCENTS[0])


prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def alt(shape, text):
    shape._element._nvXxPr.cNvPr.set("descr", text)


def rect(slide, x, y, w, h, fill=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
         line=None, radius=None, alt_text=None, transparent=False):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if transparent:
        s.fill.background()
    if fill is not None:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if alt_text:
        alt(s, alt_text)
    else:
        alt(s, "Élément graphique décoratif, sans information propre")
    return s


def textbox(slide, x, y, w, h, text, size=16, color=INK, bold=False,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.12,
            alt_text=None, italic=False, space_after=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if space_after:
            p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = ln
        f = r.font
        f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
        f.color.rgb = color
    if alt_text:
        alt(tb, alt_text)
    return tb, tf


def fit_title(slide, x, y, w, h, text, start=30, color=INK, bold=True,
              align=PP_ALIGN.LEFT, alt_text=None, anchor=MSO_ANCHOR.TOP, spacing=1.06):
    size = fit_size(text, w * 72 - 4, h * 72, start, 15, bold)
    return textbox(slide, x, y, w, h, text, size=size, color=color, bold=bold,
                   align=align, anchor=anchor, spacing=spacing, alt_text=alt_text)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def footer(slide, idx, source, part, duration):
    rect(slide, 0, 0, 0.22, SH, fill=accent(part))
    textbox(slide, 0.62, SH - 0.52, 9.4, 0.34,
            f"Source : {source}", size=9.5, color=MUTED,
            alt_text=f"Référence de source affichée en pied de diapositive : {source}")
    textbox(slide, SW - 2.6, SH - 0.52, 1.98, 0.34,
            f"{idx:02d} / {len(C.SL):02d}  ·  {duration} min",
            size=9.5, color=MUTED, align=PP_ALIGN.RIGHT,
            alt_text=f"Numéro de diapositive {idx} sur {len(C.SL)}, durée prévue {duration} minutes")


def eyebrow(slide, text, part):
    textbox(slide, 0.62, 0.42, 9.0, 0.36, text.upper(), size=12.5,
            color=accent(part), bold=True,
            alt_text=f"Repère de section affiché en haut de diapositive : {text}")


def bullet_rows(slide, nodes, accent_color, x=0.62, y=2.15, w=12.1,
                row_h=0.92, gap=0.16, size=16.5, numbered=True, alt_prefix=""):
    for i, node in enumerate(nodes):
        ry = y + i * (row_h + gap)
        rect(slide, x, ry, w, row_h, fill=PANEL,
             alt_text=f"{alt_prefix}Point {i+1} : {node}")
        rect(slide, x, ry, 0.07, row_h, fill=accent_color, shape=MSO_SHAPE.RECTANGLE,
             alt_text="Repère décoratif de section")
        if numbered:
            textbox(slide, x + 0.3, ry + row_h / 2 - 0.3, 0.9, 0.6,
                    f"{i+1:02d}", size=24, color=accent_color, bold=True,
                    anchor=MSO_ANCHOR.MIDDLE,
                    alt_text=f"Numéro d'étape {i+1}")
        tw = w - (1.55 if numbered else 0.7)
        sz = fit_size(node, tw * 72, row_h * 72 - 8, size, 12, False)
        textbox(slide, x + (1.3 if numbered else 0.35), ry + 0.06, tw, row_h - 0.12,
                node, size=sz, color=INK, anchor=MSO_ANCHOR.MIDDLE,
                alt_text=f"Texte du point {i+1} : {node}")


def cards(slide, nodes, accent_color, y=2.3, h=2.9, x0=0.62, total_w=12.1,
          size=15.5, gap=0.24, alt_prefix=""):
    n = len(nodes)
    w = (total_w - gap * (n - 1)) / n
    for i, node in enumerate(nodes):
        cx = x0 + i * (w + gap)
        rect(slide, cx, y, w, h, fill=PANEL,
             alt_text=f"{alt_prefix}Carte {i+1} : {node}")
        rect(slide, cx, y, w, 0.09, fill=accent_color, shape=MSO_SHAPE.RECTANGLE,
             alt_text="Repère décoratif de section")
        sz = fit_size(node, w * 72 - 44, h * 72 - 70, size, 11.5, False)
        textbox(slide, cx + 0.22, y + 0.42, w - 0.44, h - 0.6, node, size=sz,
                color=INK, align=PP_ALIGN.LEFT, alt_text=f"Texte de la carte {i+1} : {node}")
        textbox(slide, cx + 0.22, y + h - 0.62, 0.72, 0.4, f"{i+1:02d}",
                size=15, color=accent_color, bold=True,
                alt_text=f"Numéro {i+1}")


def band(slide, text, accent_color, y=6.15, h=0.72, color=None):
    rect(slide, 0.62, y, 12.1, h, fill=PANEL2, alt_text=f"Bandeau d'information : {text}")
    sz = fit_size(text, 11.5 * 72, h * 72 - 6, 15.5, 11.5, True)
    textbox(slide, 0.9, y + 0.05, 11.6, h - 0.1, text, size=sz,
            color=color or accent_color, bold=True, anchor=MSO_ANCHOR.MIDDLE,
            alt_text=f"Texte du bandeau : {text}")


# ---------------------------------------------------------------- slides
for idx, s in enumerate(C.SL, 1):
    slide = prs.slides.add_slide(BLANK)
    a = accent(s.part)
    rect(slide, 0, 0, SW, SH, fill=BG, shape=MSO_SHAPE.RECTANGLE,
         alt_text="Fond bleu nuit de la diapositive")
    k = s.kind

    if k == "cover":
        rect(slide, 0, 0, SW, SH, fill=BG, shape=MSO_SHAPE.RECTANGLE)
        rect(slide, 0.9, 1.5, 0.16, 3.1, fill=a, shape=MSO_SHAPE.RECTANGLE,
             alt_text="Repère décoratif vertical")
        fit_title(slide, 1.4, 1.45, 10.5, 1.5, s.title, start=72, alt_text=f"Titre : {s.title}")
        textbox(slide, 1.4, 3.05, 10.5, 0.7, s.subtitle, size=30, color=a, bold=True,
                alt_text=f"Sous-titre : {s.subtitle}")
        textbox(slide, 1.4, 3.95, 10.0, 1.2,
                "Formation bâtie sur le manuscrit illustré adapté (DOCX) disponible dans le dépôt.\n"
                "Repères évolutifs revérifiés sur les pages officielles le 28 septembre 2026.",
                size=14, color=MUTED, alt_text="Mention de source et de date de vérification")
        textbox(slide, 1.4, 5.55, 10.0, 0.5,
                "Direction artistique proposée : bleu nuit · cyan · violet — non officielle.",
                size=12, color=MUTED, italic=True,
                alt_text="Mention : palette proposée, sans charte officielle")
    elif k == "manifesto":
        eyebrow(slide, "Cadre de la formation", s.part)
        fit_title(slide, 0.62, 1.0, 12.1, 1.0, s.title, start=40,
                  alt_text=f"Titre : {s.title}")
        textbox(slide, 0.62, 1.95, 12.1, 0.5, s.subtitle, size=19, color=a, bold=True,
                alt_text=f"Sous-titre : {s.subtitle}")
        cards(slide, s.nodes, a, y=2.75, h=2.5, size=17,
              alt_prefix="Engagement de méthode. ")
    elif k == "path":
        eyebrow(slide, "Fil conducteur", s.part)
        fit_title(slide, 0.62, 1.0, 12.1, 0.9, s.title, start=40,
                  alt_text=f"Titre : {s.title}")
        w, gap = 2.78, 0.33
        for i, node in enumerate(s.nodes):
            row, col = divmod(i, 4)
            cx, cy = 0.62 + col * (w + gap), 2.16 + row * 1.76
            rect(slide, cx, cy, w, 1.3, fill=PANEL,
                 alt_text=f"Étape {i+1} du fil conducteur : {node}")
            rect(slide, cx, cy, w, 0.08, fill=a, shape=MSO_SHAPE.RECTANGLE,
                 alt_text="Repère décoratif")
            textbox(slide, cx + 0.18, cy + 0.16, 0.4, 0.3, f"{i+1:02d}",
                    size=12, color=a, bold=True,
                    alt_text=f"Numéro d'étape {i+1}")
            sz = fit_size(node, (w - 0.36) * 72, 0.75 * 72, 17, 13, True)
            textbox(slide, cx + 0.18, cy + 0.47, w - 0.36, 0.75, node, size=sz,
                    color=INK, bold=True, anchor=MSO_ANCHOR.MIDDLE,
                    alt_text=f"Texte de l'étape {i+1} : {node}")
            if col < 3:
                rect(slide, cx + w + 0.065, cy + 0.54, 0.2, 0.22, fill=a,
                     shape=MSO_SHAPE.RIGHT_ARROW,
                     alt_text=f"Flèche vers l'étape {i+2}")
        band(slide, "Ordre : ligne supérieure de gauche à droite, puis ligne inférieure.", a)
    elif k == "section":
        rect(slide, 0, 0, SW, SH, fill=BG, shape=MSO_SHAPE.RECTANGLE)
        rect(slide, 0, 0, SW, 0.28, fill=a, shape=MSO_SHAPE.RECTANGLE,
             alt_text="Bandeau supérieur de partie")
        label = s.title
        fit_title(slide, 0.9, 2.35, 9.14, 1.2, label, start=56,
                  alt_text=f"Titre de section : {label}")
        sz = fit_size(s.subtitle, 11.0 * 72, 1.6 * 72, 30, 18, True)
        textbox(slide, 0.9, 3.6, 11.0, 1.6, s.subtitle, size=sz, color=a, bold=True,
                alt_text=f"Sous-titre de section : {s.subtitle}")
    elif k == "chapter":
        eyebrow(slide, f"Partie {s.part:02d} · chapitre {s.chapter:02d}", s.part)
        fit_title(slide, 0.62, 0.85, 12.1, 0.98, s.title, start=32,
                  alt_text=f"Titre du chapitre : {s.title}")
        textbox(slide, 0.62, 1.76, 11.7, 0.64, s.subtitle, size=15.5, color=a, bold=True,
                alt_text=f"Objectif du chapitre : {s.subtitle}")
        if len(s.nodes) >= 5:
            bullet_rows(slide, s.nodes, a, y=2.42, row_h=0.65, gap=0.12, size=15.5)
        else:
            bullet_rows(slide, s.nodes, a, y=2.55, row_h=0.78, gap=0.14, size=16)
    elif k == "exercise":
        eyebrow(slide, f"Atelier · partie {s.part:02d} · chapitre {s.chapter:02d}", s.part)
        fit_title(slide, 0.62, 0.85, 12.1, 0.98, s.title, start=34,
                  alt_text=f"Titre de l'atelier : {s.title}")
        textbox(slide, 0.62, 1.76, 11.7, 0.64, s.subtitle, size=16, color=a, bold=True,
                alt_text=f"Consigne cadre : {s.subtitle}")
        rect(slide, 0.62, 2.6, 12.1, 2.4, fill=PANEL,
             alt_text=f"Consigne d'atelier : {s.exercise}")
        textbox(slide, 1.05, 2.9, 11.3, 1.9, s.exercise, size=24, color=INK, bold=True,
                anchor=MSO_ANCHOR.MIDDLE, alt_text=f"Consigne d'atelier : {s.exercise}")
        band(slide, "Aucun exemple du livre ne doit être présenté comme un témoignage réel : il est annoncé comme simulation pédagogique.", a)
    elif k == "recap":
        eyebrow(slide, f"Synthèse · partie {s.part:02d}", s.part)
        fit_title(slide, 0.62, 0.85, 12.1, 0.98, s.title, start=36,
                  alt_text=f"Titre : {s.title}")
        textbox(slide, 0.62, 1.76, 11.7, 0.64, s.subtitle, size=16, color=a, bold=True,
                alt_text=f"Sous-titre : {s.subtitle}")
        if len(s.nodes) >= 5:
            bullet_rows(slide, s.nodes, a, y=2.42, row_h=0.65, gap=0.12, size=15)
        else:
            bullet_rows(slide, s.nodes, a, y=2.5, row_h=0.72, gap=0.13, size=15.5)
    elif k == "prompt":
        eyebrow(slide, f"Bloc pratique · chapitre {s.chapter:02d}", s.part)
        fit_title(slide, 0.62, 0.85, 12.1, 0.98, s.title, start=32,
                  alt_text=f"Titre : {s.title}")
        textbox(slide, 0.62, 1.76, 11.7, 0.64, s.subtitle, size=16, color=a, bold=True,
                alt_text=f"Sous-titre : {s.subtitle}")
        cards(slide, s.nodes, a, y=2.55, h=1.85, size=15.5,
              alt_prefix="Variable à compléter. ")
        band(slide, "Version complète du prompt dans les notes du présentateur · validation humaine obligatoire.", a)
    elif k == "facts":
        eyebrow(slide, f"Vérification officielle · chapitre {s.chapter:02d}", s.part)
        fit_title(slide, 0.62, 0.85, 12.1, 0.98, s.title, start=32,
                  alt_text=f"Titre : {s.title}")
        textbox(slide, 0.62, 1.76, 11.7, 0.64, s.subtitle, size=16, color=a, bold=True,
                alt_text=f"Sous-titre : {s.subtitle}")
        if len(s.nodes) >= 5:
            bullet_rows(slide, s.nodes, a, y=2.42, row_h=0.65, gap=0.12, size=15)
        else:
            bullet_rows(slide, s.nodes, a, y=2.5, row_h=0.72, gap=0.13, size=15.5)
    elif k == "decision":
        eyebrow(slide, "Décision", s.part)
        fit_title(slide, 0.62, 0.85, 12.1, 0.98, s.title, start=34,
                  alt_text=f"Titre : {s.title}")
        textbox(slide, 0.62, 1.76, 11.7, 0.64, s.subtitle, size=16, color=a, bold=True,
                alt_text=f"Sous-titre : {s.subtitle}")
        pos = [(0.62, 2.6), (6.8, 2.6), (0.62, 4.45), (6.8, 4.45)]
        for i, node in enumerate(s.nodes):
            x, y = pos[i]
            rect(slide, x, y, 5.92, 1.7, fill=PANEL,
                 alt_text=f"Question de décision {i+1} : {node}")
            rect(slide, x, y, 0.07, 1.7, fill=a, shape=MSO_SHAPE.RECTANGLE,
                 alt_text="Repère décoratif")
            sz = fit_size(node, 5.2 * 72, 1.3 * 72, 17, 12.5, True)
            textbox(slide, x + 0.3, y + 0.2, 5.4, 1.3, node, size=sz, color=INK,
                    bold=True, anchor=MSO_ANCHOR.MIDDLE,
                    alt_text=f"Texte de la décision {i+1} : {node}")
            textbox(slide, x + 5.0, y + 1.22, 0.8, 0.34, f"{i+1:02d}", size=13,
                    color=a, bold=True, align=PP_ALIGN.RIGHT,
                    alt_text=f"Numéro {i+1}")
    elif k == "roadmap":
        eyebrow(slide, "Mise en œuvre", s.part)
        fit_title(slide, 0.62, 0.85, 12.1, 0.98, s.title, start=36,
                  alt_text=f"Titre : {s.title}")
        n = len(s.nodes)
        w = (12.1 - 0.3 * (n - 1)) / n
        for i, node in enumerate(s.nodes):
            cx = 0.62 + i * (w + 0.3)
            rect(slide, cx, 2.45, w, 2.5, fill=PANEL,
                 alt_text=f"Étape {i+1} du calendrier : {node}")
            rect(slide, cx, 2.45, w, 0.1, fill=a, shape=MSO_SHAPE.RECTANGLE,
                 alt_text="Repère décoratif")
            textbox(slide, cx + 0.2, 2.65, w - 0.4, 0.5, f"ÉTAPE {i+1}", size=12.5,
                    color=a, bold=True, alt_text=f"Étiquette étape {i+1}")
            sz = fit_size(node, w * 72 - 40, 1.5 * 72, 16.5, 11.5, False)
            textbox(slide, cx + 0.2, 3.2, w - 0.4, 1.5, node, size=sz, color=INK,
                    alt_text=f"Texte de l'étape {i+1} : {node}")
            if i < n - 1:
                rect(slide, cx + w + 0.04, 3.6, 0.22, 0.24, fill=a,
                     shape=MSO_SHAPE.RIGHT_ARROW,
                     alt_text=f"Flèche de progression vers l'étape {i+2}")
        band(slide, "Calendrier indicatif : les délais ne sont ni un seuil de réussite ni une promesse.", a)
    elif k == "checklist":
        eyebrow(slide, "Contrôle qualité", s.part)
        fit_title(slide, 0.62, 0.85, 12.1, 0.98, s.title, start=36,
                  alt_text=f"Titre : {s.title}")
        y = 2.3
        for i, node in enumerate(s.nodes):
            ry = y + i * 0.82
            rect(slide, 0.62, ry, 12.1, 0.7, fill=PANEL,
                 alt_text=f"Contrôle {i+1} : {node}")
            rect(slide, 0.95, ry + 0.16, 0.38, 0.38, fill=None, line=a,
                 shape=MSO_SHAPE.RECTANGLE, alt_text=f"Case à cocher {i+1}")
            sz = fit_size(node, 10.6 * 72, 0.6 * 72, 17, 12, False)
            textbox(slide, 1.55, ry + 0.08, 10.9, 0.54, node, size=sz, color=INK,
                    anchor=MSO_ANCHOR.MIDDLE, alt_text=f"Texte du contrôle {i+1} : {node}")
    elif k == "closing":
        rect(slide, 0, 0, SW, SH, fill=BG, shape=MSO_SHAPE.RECTANGLE)
        rect(slide, 0, 0, SW, 0.28, fill=a, shape=MSO_SHAPE.RECTANGLE,
             alt_text="Bandeau supérieur")
        fit_title(slide, 0.9, 1.15, 11.5, 1.1, s.title, start=54,
                  alt_text=f"Titre : {s.title}")
        chain = ["Une compétence", "Un public", "Un contenu pilote", "Une mesure",
                 "Une décision"]
        w = (11.54 - 0.2 * 4) / 5
        for i, node in enumerate(chain):
            cx = 0.9 + i * (w + 0.2)
            rect(slide, cx, 2.75, w, 1.5, fill=PANEL,
                 alt_text=f"Étape de clôture {i+1} : {node}")
            rect(slide, cx, 2.75, w, 0.08, fill=a, shape=MSO_SHAPE.RECTANGLE,
                 alt_text="Repère décoratif")
            sz = fit_size(node, w * 72 - 34, 100, 16, 11.5, True)
            textbox(slide, cx + 0.16, 2.85, w - 0.32, 1.3, node, size=sz, color=INK,
                    bold=True, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER,
                    alt_text=f"Texte : {node}")
        rect(slide, 0.9, 4.6, 11.54, 1.25, fill=PANEL2,
             alt_text="Phrase à compléter par chaque participant")
        textbox(slide, 1.15, 4.75, 11.0, 0.95,
                "Je peux aider [public] à [tâche] en produisant [contenu] ; je vérifierai [fait ou droit] et mesurerai [indicateur défini].",
                size=17, color=a, bold=True, anchor=MSO_ANCHOR.MIDDLE,
                alt_text="Modèle de phrase à compléter")
        textbox(slide, 0.9, 6.15, 11.5, 0.5,
                "Aucun revenu, aucune viralité et aucune monétisation ne sont garantis.",
                size=15, color=MUTED, alt_text="Avertissement : aucune garantie de revenu")
    elif k == "resources":
        eyebrow(slide, "Ressources", s.part)
        fit_title(slide, 0.62, 0.85, 12.1, 0.98, s.title, start=38,
                  alt_text=f"Titre : {s.title}")
        bullet_rows(slide, s.nodes, a, y=2.35, row_h=0.82, gap=0.16, size=16)
        band(slide, "Sources ouvertes le 28 septembre 2026 : voir BELIVE_MONEY_credits_sources_images.md", a)

    footer(slide, idx, s.slidesource, s.part, s.duration)
    add_notes(slide, f"MINUTAGE : {s.duration} min\nID : {s.id}\n\n{s.notes}")

prs.save(OUT)
print(f"OK {OUT.name} — {len(prs.slides.__iter__.__self__._sldIdLst)} diapositives")
